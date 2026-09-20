# -*- coding: utf-8 -*-
"""tools_office 功能测试：生成 → 重新打开 → 逐项核对内容。

判据（动手前定死，见本文件末尾标准输出）：
  1. docx / xlsx / pptx 三个格式生成后重新打开，内容逐项一致
  2. 生成物是合法 OOXML（zip 内含 [Content_Types].xml）
  3. fill_docx 能替换被拆到多个 run 的 {{占位符}}
  4. 路径闸：../ 、盘符、/ 开头全部被拒且报错可读
  5. 中文不乱码（写进去 == 读回来）

跑法：  cd secretary && py -X utf8 test_office.py
结果落盘：output/test_results/office_test.json 与 office_test.md
"""
import json
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import tools_office as T

RESULTS = []
SUB = 'test_results/office_test'          # 相对 output 的产物目录


def check(name, ok, detail=''):
    RESULTS.append({'case': name, 'ok': bool(ok), 'detail': str(detail)[:300]})
    print(('  [OK]   ' if ok else '  [FAIL] ') + name + (('   ' + str(detail)[:160]) if detail else ''))
    return ok


def is_ooxml(rel_path):
    """合法 OOXML：是个 zip，且根目录有 [Content_Types].xml。"""
    p = os.path.join(T.output_dir(), rel_path)
    if not zipfile.is_zipfile(p):
        return False, '不是 zip'
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
    return ('[Content_Types].xml' in names), ('entries=%d' % len(names))


def main():
    print('== 1. docx 生成 → 读回 ==')
    blocks = [
        {'type': 'heading', 'text': '2026年8月经营分析', 'level': 1},
        {'type': 'para', 'text': '本月售电量完成情况良好，同比增长 3.2%。'},
        {'type': 'table', 'header': ['站所', '预算(万元)', '执行率'],
         'rows': [['两水供电所', 1880, '92%'], ['解河供电所', 960, '78%']]},
        {'type': 'pagebreak'},
        {'type': 'para', 'text': '第二页说明'},
    ]
    r = T.make_docx(SUB + '/报告.docx', blocks)
    check('make_docx 返回路径', r.get('path') == SUB + '/报告.docx', r)
    doc = T.read_docx(r['path'])
    texts = [p['text'] for p in doc.get('paragraphs', [])]
    check('docx 标题读回一致', '2026年8月经营分析' in texts)
    check('docx 中文正文一致',
          any('同比增长 3.2%' in t for t in texts), texts[:3])
    tb = doc.get('tables') or [[]]
    check('docx 表格读回一致',
          tb[0][:2] == [['站所', '预算(万元)', '执行率'], ['两水供电所', '1880', '92%']], tb[0][:2])

    print('== 2. docx 跨 run 占位符替换（真实模板的常态） ==')
    from docx import Document
    tpl = os.path.join(T.output_dir(), SUB, '模板.docx')
    os.makedirs(os.path.dirname(tpl), exist_ok=True)
    d = Document()
    para = d.add_paragraph()
    para.add_run('供电所：{{na')      # 故意把 {{name}} 拆到两个 run
    para.add_run('me}}，统计月份 {{month}}')
    tb = d.add_table(rows=1, cols=2)
    tb.style = 'Table Grid'
    tb.rows[0].cells[0].text = '负责人'
    tb.rows[0].cells[1].text = '{{owner}}'
    d.save(tpl)
    r = T.fill_docx(SUB + '/模板.docx', {'name': '两水供电所', 'month': '2026-08', 'owner': '张三'},
                    out=SUB + '/已填.docx')
    check('fill_docx 替换数 == 3', r.get('replaced') == 3, r)
    filled = T.read_docx(r['path'])
    ft = ' '.join(p['text'] for p in filled['paragraphs'])
    check('跨 run 占位符被替换', '供电所：两水供电所，统计月份 2026-08' in ft, ft)
    check('表格内占位符被替换', filled['tables'][0][0][1] == '张三', filled['tables'][0])
    check('模板原件未被改动',
          '{{owner}}' in ' '.join(c for row in T.read_docx(SUB + '/模板.docx')['tables'][0] for c in row))
    left = T.fill_docx(SUB + '/模板.docx', {'name': 'X'}, out=SUB + '/已填2.docx')
    check('未知占位符原样保留',
          '{{month}}' in ' '.join(p['text'] for p in T.read_docx(left['path'])['paragraphs']))

    print('== 3. xlsx 生成 → 读回 ==')
    r = T.write_xlsx(SUB + '/预算表.xlsx', [
        {'name': '预算', 'header': ['站所', '预算(万元)'],
         'rows': [['两水供电所', 1880], ['解河供电所', 960]]},
        {'name': '执行', 'header': ['站所', '执行率'], 'rows': [['两水供电所', 0.92]]},
    ])
    check('write_xlsx 返回路径', r.get('path') == SUB + '/预算表.xlsx', r)
    x = T.read_xlsx(r['path'])
    check('xlsx 表名齐全', x.get('sheets') == ['预算', '执行'], x.get('sheets'))
    check('xlsx 第一张表内容一致',
          x['rows'][0] == ['站所', '预算(万元)'] and x['rows'][1][0] == '两水供电所'
          and float(x['rows'][1][1]) == 1880.0, x['rows'][:2])
    x2 = T.read_xlsx(r['path'], sheet='执行')
    check('xlsx 按表名取数', x2['rows'][1] == ['两水供电所', 0.92], x2['rows'])

    print('== 4. xlsx 就地改：保留原有内容与其他表 ==')
    r = T.update_xlsx(SUB + '/预算表.xlsx', '预算', {'A1': '站所名称', 'A4': '合计', 'B4': 2840})
    check('update_xlsx 改了 3 处', r.get('changed') == 3, r)
    x = T.read_xlsx(SUB + '/预算表.xlsx', sheet='预算')
    check('改动生效', x['rows'][0][0] == '站所名称' and x['rows'][3][:2] == ['合计', 2840.0], x['rows'])
    check('未动的行还在', x['rows'][2][0] == '解河供电所', x['rows'])
    check('另一张表没被动',
          T.read_xlsx(SUB + '/预算表.xlsx', sheet='执行')['rows'][1][0] == '两水供电所')
    T.update_xlsx(SUB + '/预算表.xlsx', '预算', {'D1': '=1+1'})
    check('公式注入被挡（按文本存）',
          T.read_xlsx(SUB + '/预算表.xlsx', sheet='预算')['rows'][0][3] == '=1+1')

    print('== 5. pptx 生成 → 重新打开 ==')
    r = T.make_pptx(SUB + '/汇报.pptx', [
        {'title': '本月完成情况', 'bullets': ['售电量同比增长 3.2%', {'text': '线损率 2.61%', 'level': 1}]},
        {'title': '分站所明细', 'table': {'header': ['站所', '执行率'], 'rows': [['两水供电所', '92%']]}},
    ], title='2026年8月经营汇报', subtitle='项目管理秘书自动生成')
    check('make_pptx 生成 3 页', r.get('slides') == 3, r)
    from pptx import Presentation
    prs = Presentation(os.path.join(T.output_dir(), r['path']))
    check('pptx 页数读回一致', len(prs.slides) == 3, len(prs.slides))
    check('pptx 封面标题一致', prs.slides[0].shapes.title.text == '2026年8月经营汇报')
    s1 = prs.slides[1]
    check('pptx 要点中文一致',
          s1.shapes.title.text == '本月完成情况'
          and '售电量同比增长 3.2%' in s1.placeholders[1].text_frame.text,
          s1.placeholders[1].text_frame.text[:60])
    s2 = prs.slides[2]
    tbl = [sh.table for sh in s2.shapes if sh.has_table][0]
    check('pptx 表格内容一致',
          tbl.cell(0, 0).text == '站所' and tbl.cell(1, 0).text == '两水供电所',
          [tbl.cell(0, 0).text, tbl.cell(1, 0).text])
    check('pptx 尺寸为 16:9',
          round(prs.slide_width / prs.slide_height, 3) == round(13.333 / 7.5, 3))

    print('== 6. OOXML 合法性 ==')
    for f in ('报告.docx', '预算表.xlsx', '汇报.pptx'):
        ok, detail = is_ooxml(SUB + '/' + f)
        check('ooxml 合法：' + f, ok, detail)

    print('== 7. 路径闸 ==')
    bads = ['../逃逸.docx', '..\\逃逸.xlsx', 'C:/Windows/x.docx', '/etc/x.docx',
            'a/../../../逃逸.docx']
    for bad in bads:
        try:
            res = T.make_docx(bad, [{'type': 'para', 'text': 'x'}])
            got = res.get('error') if isinstance(res, dict) else None
        except ValueError as e:
            got = str(e)
        check('拒绝 ' + bad, bool(got), got or '没被拒绝')
    escapes = [os.path.join(os.path.dirname(T.output_dir()), '逃逸.docx'),
               os.path.join(os.path.dirname(os.path.dirname(T.output_dir())), '逃逸.docx')]
    check('越界文件确实没落盘', not any(os.path.exists(e) for e in escapes),
          [e for e in escapes if os.path.exists(e)])

    print('== 8. 读不存在的文件：返回 error 不抛异常 ==')
    for fn, arg in (('read_docx', 'nope.docx'), ('read_xlsx', 'nope.xlsx'),
                    ('update_xlsx', 'nope.xlsx')):
        res = getattr(T, fn)(arg) if fn != 'update_xlsx' else T.update_xlsx(arg, None, {'A1': 1})
        check(fn + ' 返回 error', 'error' in res, res)

    ok_n = sum(1 for r in RESULTS if r['ok'])
    print('\n===== %d/%d 通过 =====' % (ok_n, len(RESULTS)))
    for r in RESULTS:
        if not r['ok']:
            print('FAIL: ' + r['case'] + '  ' + r['detail'])
    dump(ok_n)
    return 0 if ok_n == len(RESULTS) else 1


def dump(ok_n):
    out = os.path.join(T.output_dir(), 'test_results')
    os.makedirs(out, exist_ok=True)
    with open(os.path.join(out, 'office_test.json'), 'w', encoding='utf-8') as f:
        json.dump({'passed': ok_n, 'total': len(RESULTS), 'cases': RESULTS},
                  f, ensure_ascii=False, indent=1)
    lines = ['# tools_office 测试结果', '',
             '通过 %d / %d' % (ok_n, len(RESULTS)), '',
             '| 用例 | 结果 | 备注 |', '|---|---|---|']
    for r in RESULTS:
        lines.append('| %s | %s | %s |' % (r['case'], 'OK' if r['ok'] else 'FAIL',
                                           r['detail'].replace('|', '/')))
    lines += ['', '产物目录：output/' + SUB]
    with open(os.path.join(out, 'office_test.md'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')
    print('结果已落盘：output/test_results/office_test.json、office_test.md')


if __name__ == '__main__':
    sys.exit(main())
